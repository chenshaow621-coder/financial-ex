from __future__ import annotations

import re
from functools import lru_cache
from pathlib import Path
from typing import Any

from docx import Document

from data_loader import clean_text

PROJECT_ROOT = Path(__file__).resolve().parent.parent
RAW_DIR = PROJECT_ROOT / "data" / "raw"
PROMPT_OVERRIDE_DIR = PROJECT_ROOT / "data" / "prompt_overrides"

PROMPT_DOC_PATHS = {
    "atom_enhanced": RAW_DIR / "单条原子版-增强稿.docx",
    "atom_minimum": RAW_DIR / "原子知识最小可执行颗粒度判断提示词.docx",
    "set_closure": RAW_DIR / "合规判断主提示词.docx",
}

PLACEHOLDER_PATTERN = re.compile(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}")

PROMPT_SPECS: dict[str, dict[str, Any]] = {
    "extract_stage1": {
        "group": "抽取链路",
        "title": "Stage1 宏观类型扫描",
        "description": "用于判断条文包含哪些宏观实体分类。",
        "placeholders": ["sentence"],
        "override_path": PROMPT_OVERRIDE_DIR / "extract_stage1.txt",
        "default_text": """你是一位严谨的金融合规数据架构师。当前任务是：金融法规的命名实体识别（NER）- 阶段一。

【任务目标】
请阅读下述法律条文，判断其中包含了以下哪些宏观实体分类：
1. 法律主体(WHO)：执行主体、相关当事人、监管机构等。
2. 核心事项(WHAT)：核心对象、金融工具、法定权利义务等。
3. 时间情境(WHEN)：触发时机、法定期间、时间条件等。
4. 空间情境(WHERE)：适用场景、地理范围、系统环境等。
5. 规则类型(RULES)：强制义务、禁止、授权允许等规则定性。
6. 业务分类：支付结算、票据业务等。

【输入条文】
{{ sentence }}

【输出要求】
请仅输出条文中实际包含的分类名称，以英文逗号分隔。如果全都不包含，请输出：无。
严禁输出任何多余的解释说明或 Markdown 标记。

示例输出：法律主体(WHO),核心事项(WHAT),时间情境(WHEN)
""",
        "sort": 10,
    },
    "extract_stage2": {
        "group": "抽取链路",
        "title": "Stage2 实体精准提取",
        "description": "用于结合参考词典抽取指定类别实体。",
        "placeholders": ["target_category", "reference_dict", "sentence"],
        "override_path": PROMPT_OVERRIDE_DIR / "extract_stage2.txt",
        "default_text": """你是一位严谨的金融合规数据架构师。当前任务是：金融法规的命名实体识别（NER）- 阶段二。

【当前提取目标】
宏观分类：{{ target_category }}

【参考知识库（动态字典）】
为了统一标准，以下是该分类下我们已知的规范实体词汇（供参考）：
[{{ reference_dict }}]

【提取核心指令】
请从输入的法律条文中，提取出所有属于“{{ target_category }}”类型的实体。
1. 字典对齐：如果条文中出现了参考知识库中的词汇，请务必提取。
2. 动态扩表（极重要）：如果条文中出现了不在参考知识库中，但在语义上绝对属于“{{ target_category }}”的新实体，也必须提取出来！我们依赖你发现新概念。
3. 忠于原文：提取条文中的原始表述，不要擅自概括或缩写。

【输入条文】
{{ sentence }}

【输出要求】
请严格以 JSON 数组格式输出提取结果（确保是合法的 JSON，不要带 ```json 代码块标记）。
格式规范：
[
  {"entity_name": "提取出的词汇1", "entity_type": "{{ target_category }}"},
  {"entity_name": "提取出的词汇2", "entity_type": "{{ target_category }}"}
]
如果在条文中未找到对应实体，请严格输出空数组：[]。
严禁输出任何解释性废话。
""",
        "sort": 20,
    },
    "extract_stage3": {
        "group": "抽取链路",
        "title": "Stage3 原子规则组装",
        "description": "用于基于已抽取实体拼装 LegalAtom。",
        "placeholders": ["ner_entities_json", "sentence"],
        "override_path": PROMPT_OVERRIDE_DIR / "extract_stage3.txt",
        "default_text": """你是一位顶尖的金融合规数据架构师。当前任务是：基于已提取的实体，进行【事件组装与逻辑裂变】。

【已知情报】
我们已经从法规中提取了关键实体碎片（JSON格式）：
{{ ner_entities_json }}

【原始法规文本】
{{ sentence }}

【核心任务指令】
请你根据上述实体碎片和原始文本，将其组装为“原子规则（LegalAtom）”。

1. 核心规则分类 (Rule Type)：
   必须从以下8大类中选择：OBL_MANDATORY(强制义务), PRO_FORBIDDEN(禁止), PER_AUTH(授权允许), EVT_TRIGGER(条件触发), VAL_THRESHOLD(阈值约束), PRC_FLOW(程序流程), OBL_ONGOING(持续性义务), RPT_DISCLOSURE(披露报告)。

2. 逻辑裂变 (复合句拆分 - 极其重要)：
   - 若文本包含转折（“但是”、“但”、“除...外”）或条件分句，必须物理拆分为多个独立的原子规则。
   - 前半段一般性规定：relation_type 设为 "DEFAULT"，parent_atom_id 为 null。
   - “但是...”转折句：relation_type 设为 "EXCEPTION"，parent_atom_id 指向关联的前半段 temp_id。
   - “...的除外”排除句：relation_type 设为 "EXCLUSION"，parent_atom_id 指向关联的前半段 temp_id。

3. 人工审核护栏 (Quality Control)：
   - 扫描模糊程度词（重大过失、恶意等）、模糊时间词（及时、合理期限）、外部引用不明（另行规定且未写明出处）、缺少违反后果的禁止性条款。
   - 只要发现，必须将 `is_ambiguous` 设为 true，并对应填写 `review_reason` (AMBIGUOUS_TEXT 或 MISSING_PARAM)。

4. 4W1H 组装：
   利用已知情报中的实体，将 Who, When, Where, What, How 填充完整。如果文本中确实没有某一项，填 "未指定"。

【输出要求】
请严格输出符合以下结构的 JSON（不要带 Markdown 代码块标记，确保能被 json.loads 直接解析）：
{
  "atoms": [
    {
      "temp_id": "1",
      "rule_type": "OBL_MANDATORY",
      "behavior_struct": { "who": "...", "when": "...", "where": "...", "what": "...", "how": "..." },
      "relation_type": "DEFAULT",
      "parent_atom_id": null,
      "is_ambiguous": false,
      "review_reason": "NONE",
      "content_original": "前半段原文"
    }
  ]
}
""",
        "sort": 30,
    },
    "recall_set_closure_base": {
        "group": "召回/合规链路",
        "title": "召回闭环基础提示词",
        "description": "默认来自 data/raw/合规判断主提示词.docx。",
        "placeholders": [],
        "override_path": PROMPT_OVERRIDE_DIR / "recall_set_closure_base.txt",
        "source_path": PROMPT_DOC_PATHS["set_closure"],
        "sort": 40,
    },
    "recall_set_closure_wrapper": {
        "group": "召回/合规链路",
        "title": "召回闭环包装提示词",
        "description": "将当前问题、轮次和证据包装进召回闭环基础提示词。",
        "placeholders": ["base_prompt", "question", "round_context_json", "round_index", "evidence_json"],
        "override_path": PROMPT_OVERRIDE_DIR / "recall_set_closure_wrapper.txt",
        "default_text": """{{ base_prompt }}

[当前业务问题/审核问题]
{{ question }}

[当前命中的业务层级信息]
{{ round_context_json }}

[当前召回轮次]
第 {{ round_index }} 轮

[当前召回的知识集合]
{{ evidence_json }}

请严格按照上文要求输出 JSON，不要输出 Markdown 代码块，也不要补充额外说明。
""",
        "sort": 50,
    },
    "recall_atom_minimum_base": {
        "group": "召回/合规链路",
        "title": "原子最小颗粒度基础提示词",
        "description": "默认来自 data/raw/原子知识最小可执行颗粒度判断提示词.docx。",
        "placeholders": [],
        "override_path": PROMPT_OVERRIDE_DIR / "recall_atom_minimum_base.txt",
        "source_path": PROMPT_DOC_PATHS["atom_minimum"],
        "sort": 60,
    },
    "recall_atom_enhanced_base": {
        "group": "召回/合规链路",
        "title": "单条原子增强分析基础提示词",
        "description": "默认来自 data/raw/单条原子版-增强稿.docx。",
        "placeholders": [],
        "override_path": PROMPT_OVERRIDE_DIR / "recall_atom_enhanced_base.txt",
        "source_path": PROMPT_DOC_PATHS["atom_enhanced"],
        "sort": 70,
    },
    "recall_atom_analysis_wrapper": {
        "group": "召回/合规链路",
        "title": "单条原子增强分析包装提示词",
        "description": "将当前问题、业务上下文和原子记录包装进原子分析提示词。",
        "placeholders": ["atom_minimum_prompt", "atom_enhanced_prompt", "question", "round_context_json", "record_json"],
        "override_path": PROMPT_OVERRIDE_DIR / "recall_atom_analysis_wrapper.txt",
        "default_text": """{{ atom_minimum_prompt }}

---

{{ atom_enhanced_prompt }}

[当前业务问题]
{{ question }}

[当前命中的业务层级信息]
{{ round_context_json }}

[当前原子知识]
{{ record_json }}

请严格输出增强稿要求的 JSON，不要输出 Markdown 代码块，也不要补充额外解释。
""",
        "sort": 80,
    },
    "recall_final_conclusion": {
        "group": "召回/合规链路",
        "title": "最终结论生成提示词",
        "description": "用于从闭环后的证据生成最终合规结论。",
        "placeholders": ["allowed_conclusions", "prompt_payload_json"],
        "override_path": PROMPT_OVERRIDE_DIR / "recall_final_conclusion.txt",
        "default_text": """你是金融法规合规查验的最终结论生成器。你的任务不是继续召回，而是在证据已经基本闭环的前提下，输出稳定、克制、可解释的最终结论卡片。

要求：
1. 只能从以下结论中选择一个：{{ allowed_conclusions }}
2. 如果证据之间仍有冲突、例外条款未消解、限制性规则与授权性规则尚未完成取舍，优先输出“需人工复核”，不要强行给出“可办理/不可办理”。
3. 如果证据明确显示必须先补材料，优先输出“需补材料后办理”。
4. 输出必须是 JSON 对象，不要输出 Markdown 代码块，也不要补充额外说明。

请输出如下 JSON：
{
  "conclusion": "可办理|不可办理|有条件可办理|需补材料后办理|需人工复核|证据不足待补召回",
  "conclusion_summary": "一句话总结结论与理由",
  "confidence": 0.0,
  "legal_basis": ["最多 6 条，聚焦直接支持结论的依据"],
  "required_materials": ["最多 6 条"],
  "required_actions": ["最多 6 条"],
  "exceptions_and_limits": ["最多 8 条，包含禁止、例外、时限、阈值"],
  "missing_items": ["最多 6 条"],
  "risk_points": ["最多 6 条"],
  "follow_up_actions": ["最多 5 条，写清下一步怎么做"]
}

[输入信息]
{{ prompt_payload_json }}
""",
        "sort": 90,
    },
}


def ensure_prompt_override_dir() -> Path:
    PROMPT_OVERRIDE_DIR.mkdir(parents=True, exist_ok=True)
    return PROMPT_OVERRIDE_DIR


def load_docx_text(path: Path) -> str:
    if not path.exists():
        raise FileNotFoundError(f"Prompt doc not found: {path}")
    doc = Document(str(path))
    paragraphs = [clean_text(paragraph.text) for paragraph in doc.paragraphs if clean_text(paragraph.text)]
    return "\n".join(paragraphs)


def get_prompt_spec(prompt_key: str) -> dict[str, Any]:
    spec = PROMPT_SPECS.get(prompt_key)
    if spec is None:
        raise KeyError(f"Unknown prompt key: {prompt_key}")
    return spec


@lru_cache(maxsize=None)
def load_prompt_text(prompt_key: str) -> str:
    spec = get_prompt_spec(prompt_key)
    override_path = Path(spec["override_path"])
    if override_path.exists():
        return override_path.read_text(encoding="utf-8")
    source_path = spec.get("source_path")
    if source_path:
        return load_docx_text(Path(source_path))
    default_text = spec.get("default_text")
    if default_text is None:
        raise ValueError(f"Prompt '{prompt_key}' does not define a default_text or source_path.")
    return str(default_text)


def clear_prompt_template_cache() -> None:
    load_prompt_text.cache_clear()


def save_prompt_override(prompt_key: str, text: str) -> Path:
    spec = get_prompt_spec(prompt_key)
    content = str(text or "").replace("\r\n", "\n").strip()
    if not content:
        raise ValueError("Prompt content cannot be empty.")
    ensure_prompt_override_dir()
    override_path = Path(spec["override_path"])
    override_path.write_text(content + "\n", encoding="utf-8")
    clear_prompt_template_cache()
    return override_path


def reset_prompt_override(prompt_key: str) -> bool:
    spec = get_prompt_spec(prompt_key)
    override_path = Path(spec["override_path"])
    existed = override_path.exists()
    if existed:
        override_path.unlink()
    clear_prompt_template_cache()
    return existed


def render_prompt_template(prompt_key: str, **variables: Any) -> str:
    template = load_prompt_text(prompt_key)
    missing: list[str] = []

    def replace(match: re.Match[str]) -> str:
        name = match.group(1)
        if name not in variables:
            missing.append(name)
            return match.group(0)
        return str(variables[name])

    rendered = PLACEHOLDER_PATTERN.sub(replace, template)
    if missing:
        missing_text = ", ".join(sorted(set(missing)))
        raise KeyError(f"Prompt '{prompt_key}' is missing variables: {missing_text}")
    return rendered


def list_prompt_records() -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for key, spec in sorted(PROMPT_SPECS.items(), key=lambda item: (item[1].get("sort", 0), item[0])):
        override_path = Path(spec["override_path"])
        source_path = spec.get("source_path")
        records.append(
            {
                "key": key,
                "group": spec["group"],
                "title": spec["title"],
                "description": spec["description"],
                "placeholders": list(spec.get("placeholders", [])),
                "override_path": str(override_path),
                "source_path": str(source_path) if source_path else "",
                "is_override": override_path.exists(),
                "active_source": str(override_path if override_path.exists() else source_path or "embedded_default"),
                "text": load_prompt_text(key),
            }
        )
    return records
